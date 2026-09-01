from pathlib import Path
from typing import List

from langchain_core.documents import Document
from langchain_postgres import PGVector
from langchain_text_splitters import RecursiveCharacterTextSplitter

from pypdf import PdfReader
from docx import Document as DocxDocument

from config import settings

Path("uploads").mkdir(exist_ok=True)

vectorstore = PGVector(
    connection=settings.database_url,
    collection_name="agentic_chatbot_docs",
    embeddings=None,
)


def get_embeddings():
    from langchain_community.embeddings import FastEmbedEmbeddings
    return FastEmbedEmbeddings(
        model_name="sentence-transformers/all-MiniLM-L6-v2",
    )


def init_vectorstore():
    global vectorstore
    vectorstore = PGVector(
        connection=settings.database_url,
        collection_name="agentic_chatbot_docs",
        embeddings=get_embeddings(),
    )


def read_file_text(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
            text += "\n"
        return text

    if suffix == ".docx":
        docx_document = DocxDocument(file_path)
        return "\n".join(paragraph.text for paragraph in docx_document.paragraphs)

    if suffix in [".txt", ".md", ".py", ".csv"]:
        return path.read_text(encoding="utf-8", errors="ignore")

    raise ValueError("Unsupported file type. Upload PDF, DOCX, TXT, MD, PY, or CSV.")


def delete_source_documents(thread_id: str, source: str):
    store = _get_store()
    store.delete(
        filter={"$and": [{"thread_id": {"$eq": thread_id}}, {"source": {"$eq": source}}]}
    )


def list_thread_documents(thread_id: str) -> list[dict]:
    from psycopg import Connection

    with Connection.connect(settings.database_url, autocommit=True) as conn:
        rows = conn.execute(
            """
            SELECT cmetadata->>'source' AS source, COUNT(*)::int AS chunks
            FROM langchain_pg_embedding
            WHERE cmetadata->>'thread_id' = %s
              AND cmetadata->>'source' IS NOT NULL
            GROUP BY cmetadata->>'source'
            ORDER BY source
            """,
            (thread_id,),
        ).fetchall()

    docs = []
    legacy_prefix = f"{thread_id}_"
    for source, chunks in rows:
        filename = source[len(legacy_prefix):] if source.startswith(legacy_prefix) else source
        docs.append({"source": source, "filename": filename, "chunks": chunks})
    return docs


def delete_thread_documents(thread_id: str):
    from psycopg import Connection
    with Connection.connect(settings.database_url, autocommit=True) as conn:
        conn.execute(
            "DELETE FROM langchain_pg_embedding WHERE cmetadata->>'thread_id' = %s",
            (thread_id,),
        )
    uploads = Path("uploads")
    for legacy_file in uploads.glob(f"{thread_id}_*"):
        legacy_file.unlink(missing_ok=True)
    thread_dir = uploads / thread_id
    if thread_dir.exists():
        for file in thread_dir.iterdir():
            if file.is_file():
                file.unlink(missing_ok=True)
        try:
            thread_dir.rmdir()
        except OSError:
            pass


def _get_store():
    global vectorstore
    if vectorstore.embeddings is None:
        init_vectorstore()
    return vectorstore


def add_document_to_rag(file_path: str, thread_id: str):
    store = _get_store()

    text = read_file_text(file_path)

    if not text.strip():
        raise ValueError("No text could be extracted from this file.")

    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=150)
    chunks = splitter.split_text(text)
    source = Path(file_path).name

    delete_source_documents(thread_id, source)

    docs: List[Document] = [
        Document(page_content=chunk, metadata={"thread_id": thread_id, "source": source})
        for chunk in chunks
    ]

    store.add_documents(docs)

    return {"filename": Path(file_path).name, "chunks": len(docs)}


def retrieve_from_rag(query: str, thread_id: str, k: int = 4) -> str:
    store = _get_store()
    docs = store.similarity_search(query, k=k, filter={"thread_id": thread_id})

    if not docs:
        return "No relevant uploaded document content found."

    results = []
    for i, doc in enumerate(docs, start=1):
        source = doc.metadata.get("source", "uploaded document")
        results.append(f"[Source {i}: {source}]\n{doc.page_content}")

    return "\n\n".join(results)
